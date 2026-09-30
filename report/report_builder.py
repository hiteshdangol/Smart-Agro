# -*- coding: utf-8 -*-
"""Shared document builder for the Smart Agro BCA project report.

Follows the formatting rules of the TU BCA sample reports:
  A4 page size; margins Top 1", Bottom 1", Right 1", Left 1.25".
  Body: Times New Roman 12 pt, 1.5 line spacing, justified.
  Headings: chapter 16 pt, section 14 pt, sub-section 12 pt, bold.
  Tables/figures centered; table caption above, figure caption below, bold 12 pt.
  Front matter page numbers roman (from certificate page) starting at i;
  main content numeric starting at 1, centered at bottom.
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND_DARK = (0x1A, 0x5C, 0x3A)
BRAND_MID = (0x1F, 0x6F, 0x3F)
GRAY = (0x55, 0x55, 0x55)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PNG_DIR = os.path.join(ROOT, "report", "diagrams", "png")


class ReportBuilder:
    def __init__(self):
        self.doc = Document()
        self._setup_styles()
        self.table_count = {"front": 0}
        self.figure_count = {"front": 0}
        self.chapter = None
        self.tables = []  # list of caption strings for List of Tables
        self.figures = []  # list of caption strings for List of Figures

    # ---------------------------------------------------------------- styles
    def _setup_styles(self):
        normal = self.doc.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "Times New Roman")

        def style_heading(name, size, color, bold=True, before=12, after=6):
            st = self.doc.styles[name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.color.rgb = RGBColor(*color)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True

        style_heading("Heading 1", 16, BRAND_DARK)
        style_heading("Heading 2", 14, BRAND_MID)
        style_heading("Heading 3", 12, BRAND_MID)

    # ---------------------------------------------------------------- chapter state
    def start_chapter(self, number):
        self.chapter = number
        self.table_count = {number: 0}
        self.figure_count = {number: 0}

    def _tab_no(self):
        self.table_count[self.chapter] += 1
        return self.table_count[self.chapter]

    def _fig_no(self):
        self.figure_count[self.chapter] += 1
        return self.figure_count[self.chapter]

    # ---------------------------------------------------------------- text helpers
    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def h3(self, text):
        self.doc.add_heading(text, level=3)

    def para(self, text="", bold=False, italic=False, align=None, size=None):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size is not None:
            run.font.size = Pt(size)
        return p

    def bullets(self, items, style="List Bullet"):
        for it in items:
            self.doc.add_paragraph(it, style=style)

    def numbered(self, items):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.add_run("%d. %s" % (i, it))

    def ref(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        return p

    def caption(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        return p

    def page_break(self):
        self.doc.add_page_break()

    # ---------------------------------------------------------------- tables
    def _shade_cell(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _fmt_cell(self, cell, bold=False):
        for p in cell.paragraphs:
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(11)
                r.bold = bold

    def _spacer(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run("")
        run.font.size = Pt(4)
        return p

    def add_table(self, headers, rows, cap_title=None, header_fill="D9F0E3"):
        cap = cap_title
        if cap_title and self.chapter is not None:
            cap = "Table %d.%d: %s" % (self.chapter, self._tab_no(), cap_title)
            self.tables.append(cap)
        if cap:
            self.caption(cap)
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = ""
            run = hdr[i].paragraphs[0].add_run(h)
            run.bold = True
            self._shade_cell(hdr[i], header_fill)
            self._fmt_cell(hdr[i], bold=True)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
                self._fmt_cell(cells[i])
        self._spacer()
        return t

    # ---------------------------------------------------------------- figures
    def add_figure(self, image_path, cap_title, width_inches=3.6, max_height=3.6):
        cap = cap_title
        if cap_title and self.chapter is not None:
            cap = "Figure %d.%d: %s" % (self.chapter, self._fig_no(), cap_title)
            self.figures.append(cap)
        width = width_inches
        try:
            from PIL import Image
            w, h = Image.open(image_path).size
            if h * width / w > max_height:
                width = max_height * w / h
        except Exception:
            pass
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(image_path, width=Inches(width))
        self.caption(cap)
        self._spacer()
        return p

    def figure_placeholder(self, desc, cap_title):
        cap = cap_title
        if cap_title and self.chapter is not None:
            cap = "Figure %d.%d: %s" % (self.chapter, self._fig_no(), cap_title)
            self.figures.append(cap)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(desc)
        run.italic = True
        run.font.size = Pt(11)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "F0F5F2")
        pPr.append(shd)
        self.caption(cap)
        self._spacer()
        return p

    def code(self, code_text, cap_title=None, font_size=9):
        cap = None
        if cap_title and self.chapter is not None:
            cap = "Table %d.%d: %s" % (self.chapter, self._tab_no(), cap_title)
            self.tables.append(cap)
        if cap:
            self.caption(cap)
        for line in code_text.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(font_size)
            r = run._element.rPr.rFonts
            r.set(qn("w:eastAsia"), "Consolas")
        self._spacer()

    # ---------------------------------------------------------------- front matter
    def front_title(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(18)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)

    def add_field(self, paragraph, instr):
        run = paragraph.add_run()
        f1 = OxmlElement("w:fldChar")
        f1.set(qn("w:fldCharType"), "begin")
        f1.set(qn("w:dirty"), "true")
        run._r.append(f1)
        run = paragraph.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        run._r.append(it)
        run = paragraph.add_run()
        f2 = OxmlElement("w:fldChar")
        f2.set(qn("w:fldCharType"), "separate")
        run._r.append(f2)
        run = paragraph.add_run()
        f3 = OxmlElement("w:fldChar")
        f3.set(qn("w:fldCharType"), "end")
        run._r.append(f3)

    # ---------------------------------------------------------------- sections / page numbers
    def set_pg_num(self, section, fmt, start):
        sectPr = section._sectPr
        pg = sectPr.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sectPr.append(pg)
        pg.set(qn("w:fmt"), fmt)
        pg.set(qn("w:start"), str(start))

    def set_footer_number(self, section, show=True):
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        if show:
            self.add_field(p, " PAGE ")

    def configure_sections(self):
        for s in self.doc.sections:
            s.page_width = Cm(21)
            s.page_height = Cm(29.7)
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.right_margin = Inches(1)
            s.left_margin = Inches(1.25)
        self.set_footer_number(self.doc.sections[0], show=False)
        self.set_pg_num(self.doc.sections[1], "lowerRoman", 1)
        self.set_footer_number(self.doc.sections[1], show=True)
        self.set_pg_num(self.doc.sections[2], "decimal", 1)
        self.set_footer_number(self.doc.sections[2], show=True)

    def new_section(self):
        return self.doc.add_section(WD_SECTION_START.NEW_PAGE)

    def save(self, path):
        self.configure_sections()
        self.doc.save(path)
        return path
