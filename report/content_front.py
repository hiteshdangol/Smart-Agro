# -*- coding: utf-8 -*-
"""Front matter: cover page, recommendation, approval, acknowledgment, abstract,
table of contents, lists, and abbreviations."""
from docx.enum.text import WD_ALIGN_PARAGRAPH


def build(b):
    b.table_count = {"front": 0}
    b.figure_count = {"front": 0}

    # ------------------------------------------------------------------ cover
    for _ in range(2):
        b.doc.add_paragraph()
    b.para("SMART AGRO", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=40)
    b.para("Smart Farming System", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=22)
    b.para("Machine Learning  ·  E-Commerce  ·  Digital Payment",
           italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    for _ in range(2):
        b.doc.add_paragraph()
    b.para("A Project Report", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=18)
    for _ in range(1):
        b.doc.add_paragraph()
    b.para("Submitted to", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    b.para("Department of Computer Application", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    b.para("[Name of College]", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    b.para("[City, Nepal]", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    for _ in range(1):
        b.doc.add_paragraph()
    b.para("In partial fulfilment of the requirements for the Bachelor in Computer Application (BCA)",
           align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    for _ in range(1):
        b.doc.add_paragraph()
    b.para("Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    b.para("[Student Name]", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
    b.para("[TU Registration Number]", align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    for _ in range(1):
        b.doc.add_paragraph()
    b.para("Under the Supervision of", align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    b.para("[Supervisor Name]", align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
    for _ in range(1):
        b.doc.add_paragraph()
    b.para("July 2026", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)

    # ------------------------------------------------------------- front matter section
    b.new_section()

    b.front_title("Certificate")
    b.para("This is to certify that the project work entitled \u201cSmart Agro: A Smart Farming System\u201d "
           "submitted by [Student Name] ([TU Registration Number]) has been carried out under my "
           "supervision in partial fulfilment of the requirements for the degree of Bachelor in "
           "Computer Application (BCA) of Tribhuvan University. The project report is original and "
           "has not been submitted elsewhere for the award of any degree.")
    for _ in range(3):
        b.doc.add_paragraph()
    b.para("\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("[Supervisor Name]", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("Project Supervisor", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("BCA Department", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.doc.add_paragraph()
    b.para("Date: \u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.page_break()

    b.front_title("Supervisor\u2019s Recommendation")
    b.para("I hereby recommend that the project prepared under my supervision by [Student Name] "
           "entitled \u201cSmart Agro: A Smart Farming System\u201d in the partial fulfilment of the "
           "requirements for the degree of Bachelor in Computer Application (BCA) of Tribhuvan "
           "University is recommended for final evaluation.")
    for _ in range(4):
        b.doc.add_paragraph()
    b.para("\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("[Supervisor Name]", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("Project Supervisor", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("BCA Department", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.doc.add_paragraph()
    b.para("Date: \u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.page_break()

    b.front_title("Letter of Approval")
    b.para("This is to certify that the project prepared by [Student Name] entitled \u201cSmart Agro: "
           "A Smart Farming System\u201d in the partial fulfilment of the requirements for the degree "
           "of Bachelor in Computer Application (BCA) of Tribhuvan University has been evaluated. "
           "In our opinion, it is satisfactory in the scope and quality as a project for the "
           "required degree.")
    for _ in range(2):
        b.doc.add_paragraph()
    for label in [
        "[Supervisor Name]", "[Head of Department Name]",
        "[Internal Examiner Name]", "[External Examiner Name]",
    ]:
        b.para("\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
        b.para(label, align=WD_ALIGN_PARAGRAPH.RIGHT)
        b.doc.add_paragraph()
    b.para("Date: \u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026\u2026", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.page_break()

    b.front_title("Acknowledgement")
    b.para("I would like to express my sincere gratitude to my supervisor, [Supervisor Name], for the "
           "valuable guidance, constructive feedback, and constant encouragement provided throughout "
           "the design, development, and documentation of this project. His insights into software "
           "engineering and machine learning shaped the direction of the work presented in this report.")
    b.para("I am equally thankful to the Department of Computer Application at [Name of College] for "
           "providing the necessary laboratory facilities, learning environment, and institutional "
           "support that made this project possible. The constructive suggestions of the faculty "
           "members and the department head during the proposal and progress reviews helped refine "
           "both the scope and the quality of the system.")
    b.para("I also wish to acknowledge the open-source community, whose contributions in the form of "
           "frameworks, libraries, and documentation formed the technical foundation of this work. "
           "Finally, I am grateful to my family and friends for their patience, understanding, and "
           "moral support during the entire period of this academic endeavour.")
    b.para("Warm regards,", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.para("[Student Name]", align=WD_ALIGN_PARAGRAPH.RIGHT)
    b.page_break()

    b.front_title("Abstract")
    b.para("Agriculture remains the primary source of livelihood for a large share of the population in "
           "Nepal, yet farming decisions continue to depend largely on intuition, tradition, and local "
           "experience. Climate variability, unpredictable rainfall, soil degradation, and frequent "
           "outbreaks of pests and plant diseases further expose the limitations of conventional "
           "practice. This project presents Smart Agro, an integrated web-based smart farming system "
            "that combines machine learning advisory services and a digital agricultural marketplace "
           "a digital agricultural marketplace within a single platform. The system serves three "
           "categories of users: farmers, who access data-driven recommendations and trade; sellers, "
           "who list agricultural inputs for approval; and administrators, who manage users, products, "
           "orders, medicines, and permissions.")
    b.para("Three machine learning models are deployed as separate microservices. A Gaussian Naive "
           "Bayes classifier estimates the risk of pest incidence from temperature, humidity, rainfall, "
           "crop type, growth stage, and previous incidence. A Random Forest classifier recommends the "
           "most suitable crop together with a confidence score, the top five candidate crops, growing "
           "tips, and similar-condition matches. An EfficientNetB4 convolutional neural network "
            "recognizes plant diseases from leaf images across thirty-nine classes and, when a disease is "
            "confirmed, suggests chemical and organic medicines along with relevant shop products. The "
            "marketplace supports a shopping cart, order management, cash on delivery, and the eSewa "
           "digital payment gateway with HMAC-SHA256 signature verification.")
    b.para("The system is built on a React frontend, a Node.js and Express REST backend, and MongoDB "
           "for persistence, with role-based access control protecting administrative functions. "
           "Functional and system-level testing confirmed that all modules operate correctly and that "
            "the integrated workflows, including eSewa payment verification, function end to end. The "
            "project demonstrates that a multi-module platform combining machine learning and digital "
            "commerce can offer practical, accessible support for smallholder farming decisions "
            "in Nepal.")
    b.para("Keywords: Smart Farming, Machine Learning, Crop Recommendation, "
           "Pest Risk Prediction, Plant Disease Recognition, e-Commerce, eSewa, Role-Based Access "
           "Control, Precision Agriculture.", italic=True)
    b.page_break()

    b.front_title("Table of Contents")
    p = b.doc.add_paragraph()
    b.add_field(p, ' TOC \\o "1-3" \\h \\z \\u ')
    b.para("(In Microsoft Word, right-click the field above and select \u201cUpdate Field\u201d to "
           "generate the Table of Contents.)", italic=True, size=11)
    b.page_break()

    b.front_title("List of Tables")
    for cap in b.tables:
        b.para(cap)
    b.page_break()

    b.front_title("List of Figures")
    for cap in b.figures:
        b.para(cap)
    b.page_break()

    b.front_title("List of Abbreviations")
    b.add_table(
        ["Abbreviation", "Full Form"],
        [
            ["AI", "Artificial Intelligence"],
            ["API", "Application Programming Interface"],
            ["BCA", "Bachelor in Computer Application"],
            ["CNN", "Convolutional Neural Network"],
            ["COD", "Cash on Delivery"],
            ["CORS", "Cross-Origin Resource Sharing"],
            ["DFD", "Data Flow Diagram"],
            ["ER", "Entity Relationship"],
            ["HMAC", "Hash-Based Message Authentication Code"],
            ["HTTP", "HyperText Transfer Protocol"],
            ["HTTPS", "HyperText Transfer Protocol Secure"],
            ["ID", "Identifier"],
            ["JWT", "JSON Web Token"],
            ["MERN", "MongoDB, Express, React, Node.js"],
            ["ML", "Machine Learning"],
            ["MVC", "Model-View-Controller"],
            ["N", "Nitrogen"],
            ["P", "Phosphorus"],
            ["K", "Potassium"],
            ["NFR", "Non-Functional Requirement"],
            ["NPM", "Node Package Manager"],
            ["RBAC", "Role-Based Access Control"],
            ["REST", "Representational State Transfer"],
            ["SGD", "Stochastic Gradient Descent"],
            ["SHA", "Secure Hash Algorithm"],
            ["SPA", "Single-Page Application"],
            ["SVM", "Support Vector Machine"],
            ["TU", "Tribhuvan University"],
            ["UML", "Unified Modeling Language"],
        ],
        cap_title="Abbreviations",
    )
    b.page_break()

    # main content section (arabic page numbers)
    b.new_section()
